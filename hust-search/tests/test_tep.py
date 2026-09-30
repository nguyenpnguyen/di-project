import io
import pathlib
import sys
import tempfile
import unittest
from unittest.mock import patch

import httpx
import mongomock
from fastapi import HTTPException

sys.path.insert(0, str(pathlib.Path(__file__).parents[1] / "api"))
import main  # noqa: E402
import routes_bt  # noqa: E402
import tep_job  # noqa: E402
from boc_tach import tep  # noqa: E402


def pdf_bytes(text: str) -> bytes:
    """PDF một trang tối thiểu (font Helvetica, chỉ ASCII); pdfminer tự dựng lại xref."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out = b"%PDF-1.4\n"
    for i, o in enumerate(objs, 1):
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    return out + b"trailer\n<< /Root 1 0 R /Size 6 >>\n%%EOF\n"


def docx_bytes() -> bytes:
    import docx
    d = docx.Document()
    d.add_paragraph("Thông báo học bổng năm 2026")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Mã ngành", "IT1"
    b = io.BytesIO(); d.save(b)
    return b.getvalue()


def xlsx_bytes() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook(); wb.active.title = "Điểm chuẩn"
    wb.active.append(["Ngành", 27.5]); b = io.BytesIO(); wb.save(b)
    return b.getvalue()


def pptx_bytes() -> bytes:
    from pptx import Presentation
    p = Presentation(); s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Lễ khai giảng"
    b = io.BytesIO(); p.save(b)
    return b.getvalue()


class BongChuTest(unittest.TestCase):
    def test_pdf_co_lop_chu(self):
        r = tep.bong_chu(pdf_bytes("Thong bao hoc bong nam 2026 danh cho sinh vien"), "pdf")
        self.assertEqual((r["status"], r["needs_ocr"], r["n_pages"]), ("ok", False, 1))
        self.assertIn("hoc bong", r["text"])

    def test_pdf_khong_co_chu_bi_gan_needs_ocr_khong_ocr(self):
        r = tep.bong_chu(pdf_bytes(""), "pdf")
        self.assertEqual((r["status"], r["needs_ocr"], r["text"]), ("ok", True, ""))

    def test_docx_xlsx_pptx(self):
        self.assertIn("học bổng", tep.bong_chu(docx_bytes(), "docx")["text"])
        self.assertIn("Mã ngành IT1", tep.bong_chu(docx_bytes(), "docx")["text"])
        x = tep.bong_chu(xlsx_bytes(), "xlsx")
        self.assertIn("Điểm chuẩn", x["text"]); self.assertIn("27.5", x["text"])
        self.assertIn("khai giảng", tep.bong_chu(pptx_bytes(), "pptx")["text"])

    def test_dinh_dang_cu_unsupported_va_tep_hong_thanh_error_khong_nem_loi(self):
        self.assertEqual(tep.bong_chu(b"x", "doc")["status"], "unsupported")
        self.assertEqual(tep.bong_chu(b"x", "xyz")["status"], "unsupported")
        r = tep.bong_chu(b"khong phai docx", "docx")
        self.assertEqual(r["status"], "error"); self.assertTrue(r["error"])

    def test_nghi_sai_bang_ma_tcvn3(self):
        tcvn = "Th«ng b¸o tuyÓn sinh n¨m häc míi ®¹i häc " * 12
        self.assertTrue(tep.nghi_sai_bang_ma(tcvn))
        self.assertFalse(tep.nghi_sai_bang_ma("Thông báo tuyển sinh năm học mới của trường đại học " * 8))


def may_chu(routes: dict, goi: list):
    def h(req: httpx.Request):
        goi.append(str(req.url))
        if req.url.path == "/robots.txt":
            return httpx.Response(200, text="User-agent: *\nDisallow: /cam/\n")
        kq = routes.get(req.url.path)
        return kq if kq is not None else httpx.Response(404)
    return httpx.Client(transport=httpx.MockTransport(h), follow_redirects=True)


class TepJobTest(unittest.TestCase):
    def setUp(self):
        self.db = mongomock.MongoClient().db
        self.tmp = pathlib.Path(tempfile.mkdtemp())
        e = lambda dst, kind="document": {"_id": dst + "|1", "src": "https://hust.edu.vn/a.html", "dst": dst,
                                          "type": "href", "text": "t", "dst_kind": kind, "count": 1,
                                          "src_host": "hust.edu.vn", "dst_host": "x"}
        self.db.links.insert_many([
            e("https://hust.edu.vn/uploads/a.pdf"), e("https://svbk.hust.edu.vn/uploads/b.docx"),
            e("https://hust.edu.vn/uploads/c.doc"), e("https://drive.google.com/x.pdf", "external"),
            e("https://hust.edu.vn/cam/d.pdf"), e("https://hust.edu.vn/uploads/e.pdf")])

    def test_danh_muc_chi_lay_host_hust_va_doc_cu_la_unsupported(self):
        r = tep_job.danh_muc(self.db)
        self.assertEqual(r["total"], 5)
        self.assertIsNone(self.db.documents.find_one({"_id": "https://drive.google.com/x.pdf"}))
        self.assertEqual(self.db.documents.find_one({"_id": "https://hust.edu.vn/uploads/c.doc"})["status"], "unsupported")
        self.assertEqual(tep_job.danh_muc(self.db)["new"], 0)          # chạy lại không thêm

    def test_danh_muc_lay_ca_ban_ghi_kho_tho_khong_phai_html(self):
        rec = {"url": "https://hust.edu.vn/tai?download=1", "status": 200,
               "content_type": "application/pdf", "html_b64": None}
        tep_job.danh_muc(self.db, [rec])
        d = self.db.documents.find_one({"_id": "https://hust.edu.vn/tai?download=1"})
        self.assertEqual((d["ext"], d["status"]), ("pdf", "pending"))

    def test_tai_roi_boc_chu_theo_robots_kich_thuoc_va_loi(self):
        tep_job.danh_muc(self.db)
        goi = []
        cli = may_chu({"/uploads/a.pdf": httpx.Response(200, content=pdf_bytes("Thong bao hoc bong nam 2026 danh cho sinh vien"),
                                                          headers={"content-type": "application/pdf"}),
                       "/uploads/b.docx": httpx.Response(200, content=docx_bytes()),
                       "/uploads/e.pdf": httpx.Response(200, content=b"x" * 200_000)}, goi)
        r = tep_job.tai(self.db, self.tmp, client=cli, nhip=0, max_bytes=100_000)
        self.assertEqual((r["errors"], r["too_large"]), (1, 1))         # robots cấm d.pdf ; e.pdf quá lớn
        self.assertNotIn("https://hust.edu.vn/cam/d.pdf", goi)          # bị cấm thì KHÔNG gọi
        self.assertEqual(self.db.documents.find_one({"_id": "https://hust.edu.vn/cam/d.pdf"})["error"], "robots.txt cấm")
        self.assertEqual(len(list(self.tmp.iterdir())), 2)
        # gọi lại không tải lại tệp đã có sha1
        n_goi = len(goi)
        tep_job.tai(self.db, self.tmp, client=cli, nhip=0, max_bytes=100_000)
        self.assertEqual([g for g in goi[n_goi:] if g.endswith(".pdf")], [])

        kq = tep_job.boc_chu(self.db, self.tmp)
        self.assertEqual((kq["extracted"], kq["ok"]), (2, 2))
        a = self.db.documents.find_one({"_id": "https://hust.edu.vn/uploads/a.pdf"})
        self.assertEqual(a["status"], "ok"); self.assertIn("hoc bong", a["text"])
        b = self.db.documents.find_one({"_id": "https://svbk.hust.edu.vn/uploads/b.docx"})
        self.assertIn("học bổng", b["text"])

    def test_http_loi_ghi_status_error(self):
        tep_job.danh_muc(self.db)
        cli = may_chu({}, [])            # mọi đường dẫn 404
        tep_job.tai(self.db, self.tmp, client=cli, nhip=0)
        self.assertEqual(self.db.documents.find_one({"_id": "https://hust.edu.vn/uploads/a.pdf"})["error"], "HTTP 404")


class FilesRoutesTest(unittest.TestCase):
    def setUp(self):
        self.db = mongomock.MongoClient().db
        self.db.documents.insert_one({"_id": "https://hust.edu.vn/uploads/a.pdf", "host": "hust.edu.vn",
                                      "ext": "pdf", "status": "ok", "text": "nội dung"})
        self.db.links.insert_one({"_id": "1", "src": "https://hust.edu.vn/b.html", "dst": "https://hust.edu.vn/uploads/a.pdf",
                                  "type": "href", "text": "Tải về", "dst_kind": "document", "count": 1,
                                  "src_host": "h", "dst_host": "h"})
        self.db.images.insert_one({"_id": "https://hust.edu.vn/a.jpg", "host": "hust.edu.vn", "alts": ["Ảnh"], "is_template": False})
        p = patch.object(routes_bt, "mdb", lambda: self.db); p.start(); self.addCleanup(p.stop)

    def test_files_list_co_so_trang_gioi_thieu_va_khong_lo_text(self):
        r = routes_bt.files_list()
        self.assertEqual(r["items"][0]["referrers"], 1)
        self.assertNotIn("text", r["items"][0])
        self.assertEqual(r["by_status"], {"ok": 1})

    def test_images_list(self):
        self.assertEqual(routes_bt.images_list(template=False)["items"][0]["alts"], ["Ảnh"])
        self.assertEqual(routes_bt.images_list(template=True)["total"], 0)

    def test_files_fetch_tu_choi_khi_crawler_dang_chay(self):
        with patch.object(main, "_alive", lambda: object()):
            with self.assertRaises(HTTPException) as c:
                routes_bt.files_fetch()
        self.assertEqual(c.exception.status_code, 409)

    def test_crawl_start_tu_choi_khi_dang_tai_tep(self):
        routes_bt._job.update(running=True, what="files")
        try:
            with self.assertRaises(HTTPException) as c:
                main.crawl_start(main.CrawlReq())
            self.assertEqual(c.exception.status_code, 409)
        finally:
            routes_bt._job.update(running=False, what="")


if __name__ == "__main__":
    unittest.main()
