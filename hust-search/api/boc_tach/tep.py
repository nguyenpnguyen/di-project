"""Bóc chữ từ tệp tài liệu: pdf, docx, xlsx, pptx.

Không OCR: PDF scan không có lớp chữ chỉ được gắn cờ `needs_ocr`. Định dạng cũ
(doc/xls/ppt) không hỗ trợ — LibreOffice headless làm image nặng thêm nhiều.
"""
from __future__ import annotations

import io
import re

HO_TRO = {"pdf", "docx", "xlsx", "pptx"}
CU = {"doc", "xls", "ppt"}
MIME_SANG_DUOI = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/msword": "doc",
    "application/vnd.ms-excel": "xls",
    "application/vnd.ms-powerpoint": "ppt",
}
TOI_DA_KY_TU = 500_000
# PDF có ít hơn chừng này ký tự mỗi trang thì coi là scan (không có lớp chữ)
NGUONG_KY_TU_TRANG = 20


def duoi_tu_url(url: str) -> str:
    ten = url.split("?")[0].rsplit("/", 1)[-1]
    return ten.rsplit(".", 1)[-1].lower() if "." in ten else ""


def nghi_sai_bang_ma(text: str) -> bool:
    """Heuristic: văn bản tiếng Việt bảng mã cũ (TCVN3/VNI) đọc bằng Unicode ra nhiều ký tự
    Latin-1 lạ (¸ µ ¶ · ¹ ...) mà gần như không có chữ Việt có dấu. Có thể báo nhầm/sót."""
    if len(text) < 200:
        return False
    la = sum(1 for c in text if "¡" <= c <= "ÿ")
    viet = sum(1 for c in text if c in "ăâđêôơưĂÂĐÊÔƠƯ" or "Ạ" <= c <= "ỹ")
    return la / len(text) > 0.04 and viet / len(text) < 0.005


def _pdf(data: bytes) -> tuple[str, int]:
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LTTextContainer
    parts, n = [], 0
    for page in extract_pages(io.BytesIO(data)):
        n += 1
        parts.append(" ".join(e.get_text() for e in page if isinstance(e, LTTextContainer)))
    return "\n".join(parts), n


def _docx(data: bytes) -> tuple[str, int]:
    import docx
    d = docx.Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(" ".join(c.text for c in row.cells))
    return "\n".join(out), 0


def _xlsx(data: bytes) -> tuple[str, int]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(ws.title)
        for row in ws.iter_rows(values_only=True):
            out.append(" ".join(str(c) for c in row if c is not None))
    return "\n".join(out), len(wb.worksheets)


def _pptx(data: bytes) -> tuple[str, int]:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    out = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
    return "\n".join(out), len(prs.slides)


_BOC = {"pdf": _pdf, "docx": _docx, "xlsx": _xlsx, "pptx": _pptx}


def bong_chu(data: bytes, ext: str) -> dict:
    """-> {status, text, n_pages, needs_ocr, encoding_suspect, error}."""
    kq = {"status": "ok", "text": "", "n_pages": 0, "needs_ocr": False,
          "encoding_suspect": False, "error": ""}
    ext = ext.lower()
    if ext in CU or ext not in HO_TRO:
        kq["status"] = "unsupported"
        return kq
    try:
        text, n = _BOC[ext](data)
    except Exception as e:                       # tệp hỏng: ghi lỗi, không làm hỏng cả mẻ
        kq.update(status="error", error=f"{type(e).__name__}: {e}"[:300])
        return kq
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\s*\n\s*", "\n", text).strip()[:TOI_DA_KY_TU]
    kq["n_pages"] = n
    kq["text"] = text
    if ext == "pdf" and len(text) < NGUONG_KY_TU_TRANG * max(n, 1):
        kq["needs_ocr"] = True
        kq["text"] = ""                          # vài ký tự lẻ không đáng index
    kq["encoding_suspect"] = nghi_sai_bang_ma(kq["text"])
    return kq
